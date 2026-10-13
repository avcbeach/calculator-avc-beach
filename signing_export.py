"""
signing_export.py
-----------------
Builds the "Signature sheet" Excel file (layout copied from the FIVB VIS
BeachTeamSignaturesSheetReport): one block per team - two players with
birthdate / AVC pts / FIVB pts / shirt number + coach rows ("C") - and a
signature column for everyone.

Pure openpyxl, no Streamlit, so it can be tested on its own.
"""

from dataclasses import dataclass, field
from datetime import date, datetime
from typing import List, Optional

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.pagebreak import Break

FONT = "Arial"
THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BLACK_FILL = PatternFill("solid", fgColor="000000")

ROWS_PER_PAGE = 18          # data rows that fit on one landscape A4 page
LAST_COL = 11               # A..K


@dataclass
class SigPerson:
    name: str = ""
    birthdate: str = ""     # already formatted, e.g. 05-Jan-2003 (or blank)
    avc: int = 0
    fivb: int = 0


@dataclass
class SigTeam:
    nf: str
    team_name: str
    players: List[SigPerson]                       # normally 2
    coaches: List[SigPerson] = field(default_factory=list)
    total: int = 0                                 # only used for ordering


def format_birthdate(raw) -> str:
    """VIS gives '2003-01-05' / '2003-01-05T00:00:00' - show 05-Jan-2003."""
    if not raw:
        return ""
    if isinstance(raw, (date, datetime)):
        return raw.strftime("%d-%b-%Y")
    try:
        return datetime.strptime(str(raw)[:10], "%Y-%m-%d").strftime("%d-%b-%Y")
    except ValueError:
        return str(raw)


def date_range_label(start: Optional[date], end: Optional[date]) -> str:
    """18 - 21-Jun-2026 (same month) or 30-May - 02-Jun-2026."""
    if not start and not end:
        return ""
    if start and end:
        if start == end:
            return end.strftime("%d-%b-%Y")
        if (start.year, start.month) == (end.year, end.month):
            return f"{start.day:02d} - {end.strftime('%d-%b-%Y')}"
        return f"{start.strftime('%d-%b')} - {end.strftime('%d-%b-%Y')}"
    d = start or end
    return d.strftime("%d-%b-%Y")


def assign_country_codes(teams: List[SigTeam]) -> List[tuple]:
    """Sort NF A-Z, then by team total desc; the code is just the NF (AUS, AUS, ...)."""
    ordered = sorted(teams, key=lambda t: (t.nf or "~", -t.total, t.team_name))
    return [(t.nf, t) for t in ordered]


def _cell(ws, row, col, value=None, *, bold=False, size=9, h="left", border=True,
          wrap=True, color="000000"):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name=FONT, size=size, bold=bold, color=color)
    c.alignment = Alignment(horizontal=h, vertical="center", wrap_text=wrap)
    if border:
        c.border = BOX
    return c


def _box_range(ws, row1, col1, row2, col2):
    for r in range(row1, row2 + 1):
        for c in range(col1, col2 + 1):
            ws.cell(row=r, column=c).border = BOX


def build_signing_workbook(event_title: str, subtitle_line: str, form_label: str,
                           teams: List[SigTeam], output_path: str,
                           logo_path: Optional[str] = None,
                           footer_note: str = ("With the signature, the player accepts all "
                                               "the points, seed and uniform numbers.")):
    wb = Workbook()
    ws = wb.active
    ws.title = "Signature sheet"
    ws.sheet_view.showGridLines = False

    widths = {"A": 8, "B": 13.5, "C": 9, "D": 31.7, "E": 12.5, "F": 6.5, "G": 6.5,
              "H": 12.6, "I": 7.5, "J": 22, "K": 21}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v

    # ---- header block (rows 1-8) ------------------------------------------
    for r, h in {1: 8, 2: 22, 3: 5, 4: 17, 5: 17, 6: 8, 7: 5, 8: 34}.items():
        ws.row_dimensions[r].height = h

    ws.merge_cells("C2:K2")
    t = _cell(ws, 2, 3, event_title, bold=True, size=13, border=False, wrap=False)
    for col in range(3, LAST_COL + 1):
        ws.cell(row=2, column=col).border = Border(bottom=THIN)
    ws.merge_cells("C4:K4")
    _cell(ws, 4, 3, subtitle_line, size=10, border=False, wrap=False)
    ws.merge_cells("C5:K5")
    _cell(ws, 5, 3, f" {form_label}   ", size=10, border=False, wrap=False)

    if logo_path:
        try:
            img = XLImage(logo_path)
            ratio = img.width / img.height if img.height else 1
            img.height = 82
            img.width = int(82 * ratio)
            ws.add_image(img, "A1")
        except Exception:
            pass   # a bad logo file must never block the export

    heads = {1: "Country\ncode", 2: "Team name", 4: "Players / Coaches", 5: "Birthdate",
             6: "AVC\npts", 7: "FIVB\npts", 8: "Total pts", 9: "Shirt\nno.", 10: "Signatures"}
    for col in range(1, LAST_COL + 1):
        c = _cell(ws, 8, col, heads.get(col), bold=True, color="FFFFFF",
                  h="center" if col in (1, 6, 7, 8, 9) else "left")
        c.fill = BLACK_FILL
    ws.merge_cells("B8:C8")
    ws.merge_cells("J8:K8")

    # ---- team blocks -------------------------------------------------------
    row = 9
    rows_on_page = 0
    page_breaks = []
    for code, team in assign_country_codes(teams):
        coaches = team.coaches or [SigPerson()]           # always one 'C' slot
        n_rows = len(team.players) + len(coaches)
        if rows_on_page and rows_on_page + n_rows > ROWS_PER_PAGE:
            page_breaks.append(row - 1)
            rows_on_page = 0
        first, last = row, row + n_rows - 1

        for r in range(first, last + 1):
            ws.row_dimensions[r].height = 22
            ws.merge_cells(start_row=r, start_column=10, end_row=r, end_column=11)
            _box_range(ws, r, 1, r, LAST_COL)

        # merged team-level cells
        def merged(col1, col2, value, h, formula=False):
            ws.merge_cells(start_row=first, start_column=col1, end_row=last, end_column=col2)
            _cell(ws, first, col1, value, h=h)
            _box_range(ws, first, col1, last, col2)

        merged(1, 1, code, "center")
        merged(2, 3, team.team_name, "left")
        p_rows = list(range(first, first + len(team.players)))
        if len(team.players) == 2:
            p1, p2 = p_rows
            total = f"=F{p1}+G{p1}+F{p2}+G{p2}"
        else:
            total = "=" + "+".join(f"F{r}+G{r}" for r in p_rows)
        merged(8, 8, total, "center")

        r = first
        for i, p in enumerate(team.players, start=1):
            _cell(ws, r, 4, p.name)
            _cell(ws, r, 5, p.birthdate)
            _cell(ws, r, 6, p.avc, h="right")
            _cell(ws, r, 7, p.fivb, h="right")
            _cell(ws, r, 9, str(i), h="center")
            r += 1
        for p in coaches:
            _cell(ws, r, 4, p.name)
            _cell(ws, r, 5, p.birthdate)
            ws.cell(row=r, column=6).fill = BLACK_FILL
            ws.cell(row=r, column=7).fill = BLACK_FILL
            _cell(ws, r, 9, "C", h="center")
            r += 1

        row = last + 1
        rows_on_page += n_rows

    # ---- footnote ----------------------------------------------------------
    ws.row_dimensions[row].height = 24
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=LAST_COL)
    _cell(ws, row, 1, footer_note, bold=True, size=10, h="center", border=False)

    # ---- print setup -------------------------------------------------------
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:8"
    ws.print_area = f"A1:K{row}"
    ws.page_margins.left = ws.page_margins.right = 0.39
    ws.page_margins.top = ws.page_margins.bottom = 0.39
    ws.print_options.horizontalCentered = True
    for b in page_breaks:
        ws.row_breaks.append(Break(id=b))
    ws.oddFooter.left.text = f"Generated {datetime.now():%d-%b-%Y  %H:%M}"
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "Page &P of &N"
    ws.oddFooter.right.size = 8

    wb.save(output_path)
    return output_path
