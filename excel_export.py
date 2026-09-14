"""
excel_export.py
----------------
Writes a workbook that matches the structure of the AVC template you
provided (2026_AVC_Beach_Tour_Yunyang_Open... .xlsx) - black & white,
bold headings + borders only, plus an optional logo:

  Sheet "Entry List": Main Draw + Reserve tables with Pos/NF/Team/Player
                        AVC & FIVB points/Team Total, sorted by total desc.
  One "Breakdown" sheet PER FEDERATION: per-player "AVC Side - Best N" and
                        "FIVB Side - Best N" tables with Season/Date/Event
                        Type/Tournament/Teammate/Rank/Points + a total row.
"""

from dataclasses import dataclass
from typing import List, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage

from scoring_engine import PlayerScore, SelectedResult

# ---------------------------------------------------------------------
# Style constants - black & white only, no fills
# ---------------------------------------------------------------------
FONT_NAME = "Arial"
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=16)
SUBTITLE_FONT = Font(name=FONT_NAME, size=11)
SECTION_FONT = Font(name=FONT_NAME, bold=True, size=12)
HEADER_FONT = Font(name=FONT_NAME, bold=True, size=10)
BODY_FONT = Font(name=FONT_NAME, size=10)
BOLD_BODY_FONT = Font(name=FONT_NAME, bold=True, size=10)
NAME_FONT = Font(name=FONT_NAME, bold=True, size=12)

THIN = Side(style="thin", color="000000")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOTTOM_THICK = Border(bottom=Side(style="medium", color="000000"))
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")


def _bucket_label(bucket: str, tournament_name: str) -> str:
    name = (tournament_name or "").lower()
    if bucket == "pro_tour":
        if "elite" in name:
            return "Pro Tour Elite16"
        if "challenge" in name:
            return "Pro Tour Challenge"
        if "futures" in name:
            return "Pro Tour Futures"
        if "final" in name:
            return "Pro Tour Finals"
        return "Pro Tour"
    return {
        "national_tour": "National Tour",
        "world_championship": "World Championship",
        "olympics": "Olympic Games",
        "multisport": "Multiple Sports Games",
        "continental_tour": "Continental Tour",
        "zonal": "Zonal / Continental Cup",
        "championship": "Continental Championship",
        "underage": "Under-age Championship",
    }.get(bucket, bucket)


@dataclass
class TeamRow:
    pos: int
    nf: str
    team_name: str
    player1_name: str
    player1_avc: int
    player1_fivb: int
    player2_name: str
    player2_avc: int
    player2_fivb: int

    @property
    def total(self) -> int:
        return self.player1_avc + self.player1_fivb + self.player2_avc + self.player2_fivb


@dataclass
class PlayerBreakdown:
    player_name: str
    federation_code: str
    score: PlayerScore
    max_events: int = 4  # for the "Best N" label


def _merge_and_style(ws, row, start_col, end_col, value, font, align=CENTER, height=None, border=False):
    ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
    cell = ws.cell(row=row, column=start_col, value=value)
    cell.font = font
    cell.alignment = align
    if border:
        for col in range(start_col, end_col + 1):
            ws.cell(row=row, column=col).border = CELL_BORDER
    if height:
        ws.row_dimensions[row].height = height
    return cell


def _style_header_row(ws, row, headers):
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = CELL_BORDER
    ws.row_dimensions[row].height = 30


def _style_data_row(ws, row, values, ncols, bold_last=False):
    for col in range(1, ncols + 1):
        val = values[col - 1] if col - 1 < len(values) else None
        c = ws.cell(row=row, column=col, value=val)
        c.font = BOLD_BODY_FONT if (bold_last and col == ncols) else BODY_FONT
        c.border = CELL_BORDER
        c.alignment = CENTER if col == 1 or isinstance(val, (int, float)) else LEFT


def _write_entry_list_section(ws, start_row: int, title: str, teams: List[TeamRow]) -> int:
    headers = ["Pos.", "NF", "Team Name", "Player 1", "Player 1 FIVB Points",
               "Player 1 AVC Points", "Player 2", "Player 2 FIVB Points",
               "Player 2 AVC Points", "Team Total Points"]
    ncols = len(headers)

    _merge_and_style(ws, start_row, 1, ncols, f"{title} ({len(teams)} Teams)",
                      SECTION_FONT, align=LEFT, height=22)
    header_row = start_row + 1
    _style_header_row(ws, header_row, headers)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1).coordinate

    teams_sorted = sorted(teams, key=lambda t: t.total, reverse=True)
    r = header_row + 1
    for i, t in enumerate(teams_sorted, start=1):
        row_vals = [i, t.nf, t.team_name, t.player1_name, t.player1_fivb, t.player1_avc,
                    t.player2_name, t.player2_fivb, t.player2_avc, t.total]
        _style_data_row(ws, r, row_vals, ncols, bold_last=True)
        r += 1
    return r + 1  # next free row (with one blank line gap)


def write_entry_list_sheet(ws, event_title: str, event_subtitle: str, gender_label: str,
                            main_draw: List[TeamRow], reserve: List[TeamRow],
                            logo_path: Optional[str] = None):
    ncols = 10
    top_row = 1
    if logo_path:
        try:
            img = XLImage(logo_path)
            img.height = 60
            img.width = 60
            ws.add_image(img, "A1")
            ws.row_dimensions[1].height = 46
            ws.row_dimensions[2].height = 46
        except Exception:
            pass  # bad/unsupported image file - just skip the logo, don't fail the export

    _merge_and_style(ws, top_row, 1, ncols, event_title, TITLE_FONT, height=28)
    _merge_and_style(ws, top_row + 1, 1, ncols, event_subtitle, SUBTITLE_FONT, height=20)
    _merge_and_style(ws, top_row + 2, 1, ncols, f"Confirmed Entry List - {gender_label}'s",
                      SUBTITLE_FONT, height=20)

    next_row = _write_entry_list_section(ws, top_row + 4, "Main Draw", main_draw)
    if reserve:
        _write_entry_list_section(ws, next_row, "Reserve", reserve)

    widths = {"A": 6, "B": 8, "C": 24, "D": 22, "E": 12, "F": 12, "G": 22, "H": 12, "I": 12, "J": 14}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def _write_side_block(ws, row: int, side_label: str, max_events: int,
                       results: List[SelectedResult], total: int, ncols: int = 7) -> int:
    _merge_and_style(ws, row, 1, ncols, f"{side_label} - Best {max_events}",
                      BOLD_BODY_FONT, align=LEFT, height=18)
    row += 1
    headers = ["Season", "Date", "Event Type", "Tournament", "Teammate", "Rank", "Points"]
    _style_header_row(ws, row, headers)
    row += 1
    for r in results:
        row_vals = [r.season, r.end_date, _bucket_label(r.bucket, r.tournament_name),
                    r.tournament_name, r.teammate, r.rank, r.points]
        _style_data_row(ws, row, row_vals, ncols)
        row += 1
    c = ws.cell(row=row, column=ncols - 1, value="Total")
    c.font = BOLD_BODY_FONT
    c.alignment = Alignment(horizontal="right")
    tc = ws.cell(row=row, column=ncols, value=total)
    tc.font = BOLD_BODY_FONT
    tc.border = BOTTOM_THICK
    row += 2  # blank line after total
    return row


def write_breakdown_sheet(ws, players: List[PlayerBreakdown]):
    ncols = 7
    row = 1
    for pb in players:
        _merge_and_style(ws, row, 1, ncols, f"{pb.player_name} ({pb.federation_code})",
                          NAME_FONT, align=LEFT, height=22)
        row += 1
        row = _write_side_block(ws, row, "FIVB Side", pb.max_events,
                                 pb.score.fivb_selected, pb.score.fivb_total, ncols)
        row = _write_side_block(ws, row, "AVC Side", pb.max_events,
                                 pb.score.avc_selected, pb.score.avc_total, ncols)
        row += 1  # extra gap between players

    for col_letter, width in [("A", 10), ("B", 12), ("C", 22), ("D", 42), ("E", 22), ("F", 8), ("G", 10)]:
        ws.column_dimensions[col_letter].width = width


def _safe_sheet_name(name: str, used: set) -> str:
    """Excel sheet names: <=31 chars, no []:*?/\\, must be unique."""
    for ch in "[]:*?/\\":
        name = name.replace(ch, "")
    name = (name or "Unknown")[:31]
    base, n = name, 2
    while name in used:
        suffix = f" ({n})"
        name = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(name)
    return name


def build_workbook(event_title: str, event_subtitle: str, gender_label: str,
                    main_draw: List[TeamRow], reserve: List[TeamRow],
                    players: List[PlayerBreakdown], output_path: str,
                    logo_path: Optional[str] = None):
    wb = Workbook()
    ws_entry = wb.active
    ws_entry.title = "Entry List"
    write_entry_list_sheet(ws_entry, event_title, event_subtitle, gender_label,
                            main_draw, reserve, logo_path=logo_path)

    # One Breakdown tab per federation, in NF-code alphabetical order, instead
    # of one long combined sheet.
    by_federation: dict = {}
    for pb in players:
        by_federation.setdefault(pb.federation_code or "Unknown", []).append(pb)

    used_sheet_names = {"Entry List"}
    for fed_code in sorted(by_federation.keys()):
        sheet_name = _safe_sheet_name(f"{fed_code} Breakdown", used_sheet_names)
        ws = wb.create_sheet(sheet_name)
        write_breakdown_sheet(ws, by_federation[fed_code])

    wb.save(output_path)
